import io
import re
import datetime
import unicodedata
import os
import json
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Tuple

import pandas as pd
import streamlit as st
from rapidfuzz import fuzz, process
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

try:
    import phonenumbers
    from phonenumbers.phonenumberutil import NumberParseException
except Exception:
    phonenumbers = None
    NumberParseException = Exception

try:
    import psycopg2
    from psycopg2.extras import RealDictCursor
except Exception:
    psycopg2 = None

try:
    from openai import OpenAI
except Exception:
    OpenAI = None

st.set_page_config(page_title="Wholesale Lead Quality Checker V2", page_icon="✅", layout="wide")

HEADER_YELLOW = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
CELL_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
CELL_BLUE = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
CELL_AMBER = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
CELL_RED = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
DASH_CHARS = "\u2010\u2011\u2012\u2013\u2014\u2212"

# --- DATABASE & CACHE SETUP ---
def get_db_connection():
    db_url = os.environ.get("DATABASE_URL")
    if not db_url:
        try:
            db_url = st.secrets.get("DATABASE_URL", None)
        except Exception:
            db_url = None
            
    if not db_url or not psycopg2:
        return None
    try:
        conn = psycopg2.connect(db_url)
        return conn
    except Exception:
        return None

def init_db():
    conn = get_db_connection()
    if not conn:
        return
    try:
        with conn.cursor() as cur:
            cur.execute("""
                CREATE TABLE IF NOT EXISTS company_cache (
                    id SERIAL PRIMARY KEY,
                    normalized_company TEXT,
                    domain TEXT,
                    status TEXT,
                    explanation TEXT,
                    source_links TEXT,
                    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    UNIQUE(normalized_company, domain)
                );
            """)
            conn.commit()
    except Exception:
        pass
    finally:
        conn.close()

init_db()

def fetch_from_cache(company_name: str, domain: str) -> Tuple[str, str, List[str]]:
    conn = get_db_connection()
    if not conn:
        return "", "", []
    try:
        norm_c = unicodedata.normalize("NFKC", str(company_name)).strip().casefold()
        clean_d = str(domain).strip().casefold()
        with conn.cursor() as cur:
            cur.execute(
                "SELECT status, explanation, source_links FROM company_cache WHERE normalized_company = %s AND domain = %s;",
                (norm_c, clean_d)
            )
            row = cur.fetchone()
            if row:
                status, explanation, links_json = row
                links = json.loads(links_json) if links_json else []
                return status, explanation, links
    except Exception:
        pass
    finally:
        conn.close()
    return "", "", []

def save_to_cache(company_name: str, domain: str, status: str, explanation: str, links: List[str]):
    conn = get_db_connection()
    if not conn:
        return
    try:
        norm_c = unicodedata.normalize("NFKC", str(company_name)).strip().casefold()
        clean_d = str(domain).strip().casefold()
        links_json = json.dumps(links)
        with conn.cursor() as cur:
            cur.execute("""
                INSERT INTO company_cache (normalized_company, domain, status, explanation, source_links, updated_at)
                VALUES (%s, %s, %s, %s, %s, CURRENT_TIMESTAMP)
                ON CONFLICT (normalized_company, domain) 
                DO UPDATE SET status = EXCLUDED.status, explanation = EXCLUDED.explanation, 
                              source_links = EXCLUDED.source_links, updated_at = CURRENT_TIMESTAMP;
            """, (norm_c, clean_d, status, explanation, links_json))
            conn.commit()
    except Exception:
        pass
    finally:
        conn.close()


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    text = str(value)
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\u00A0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text.casefold()


def norm_key(value: Any) -> str:
    text = norm_text(value)
    text = text.replace("&", " and ")
    for ch in DASH_CHARS:
        text = text.replace(ch, " ")
    text = re.sub(r"[\\/]", " ", text)
    text = text.replace("’", "'")
    text = re.sub(r"[^a-z0-9+#.\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def clean_header(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\u00A0", " ")
    text = text.strip().lower().replace("_", " ")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


PLACEHOLDER_VALUES = {"", "picklist", "leave blank", "blank", "integer", "text", "date", "dd mm yyyy", "mm dd yyyy", "https www", "http www", "do not map"}


def is_placeholder(value: Any) -> bool:
    key = norm_key(value)
    raw = norm_text(value)
    if key in PLACEHOLDER_VALUES:
        return True
    return raw.startswith(("all accepted", "target the below", "no proof", "no toll", "no fee", "please map"))


def looks_like_template_row(row: pd.Series) -> bool:
    values = [str(v).strip() for v in row.tolist() if str(v).strip()]
    if not values:
        return True
    joined = " | ".join(values[:40]).casefold()
    markers = ["picklist", "leave blank", "integer", "text", "dd/mm/yyyy", "mm/dd/yyyy", "https://", "no toll", "fee phone", "target the below", "all accepted", "do not map", "please map"]
    if sum(1 for marker in markers if marker in joined) >= 2:
        return True
    placeholder_count = sum(1 for value in values if is_placeholder(value))
    return bool(values and placeholder_count / len(values) >= 0.6)


SYNONYM_GROUPS = [
    ["us", "usa", "u s", "u s a", "united states", "united states of america", "america"],
    ["uk", "u k", "gb", "gbr", "great britain", "united kingdom", "england", "britain"],
    ["germany", "de", "deu", "deutschland"],
    ["it", "information technology", "technology"],
    ["it operations", "information technology operations", "it ops", "technology operations"],
    ["hr", "human resources", "people"],
    ["vp", "vice president"],
    ["svp", "senior vice president"],
    ["evp", "executive vice president"],
    ["c level", "c-level", "c suite", "c-suite", "chief"],
    ["ceo", "chief executive officer"],
    ["cfo", "chief financial officer"],
    ["cio", "chief information officer"],
    ["cto", "chief technology officer"],
    ["coo", "chief operating officer"],
    ["ciso", "chief information security officer"],
    ["biz dev", "business development", "bd"],
    ["ops", "operations"],
    ["manufacturing", "manufacturing and process industries", "manufacturing process industries"],
]


def canonical_key(value: Any) -> str:
    key = norm_key(value)
    if not key:
        return ""
    for group in SYNONYM_GROUPS:
        group_keys = [norm_key(item) for item in group]
        if key in group_keys:
            return group_keys[0]
    return key


def expanded_keys(value: Any) -> List[str]:
    key = norm_key(value)
    canonical = canonical_key(value)
    output: List[str] = []
    for item in [key, canonical]:
        if item and item not in output:
            output.append(item)
    for group in SYNONYM_GROUPS:
        group_keys = [norm_key(item) for item in group]
        if key in group_keys or canonical in group_keys:
            for item in group_keys:
                if item and item not in output:
                    output.append(item)
    return output


FIELD_ALIASES: Dict[str, List[str]] = {
    "company": ["company", "company name", "account", "account name", "organisation", "organization", "business name", "companyname"],
    "email": ["email", "email address", "work email", "contact email", "person email"],
    "country": ["country", "lead country", "company country", "country code", "cntry"],
    "region": ["region", "state", "province", "c state", "company state", "lead state"],
    "industry": ["industry", "companyindustry", "company industry", "main industry", "indcode1", "indcode1 main industry", "c industry"],
    "sub_industry": ["sub industry", "subindustry", "indcode2", "indcode2 sub industry"],
    "function": ["function", "department", "departments", "job function", "job area", "ocpcode1", "ocpcode1 job area", "ocpcode2", "ocpcode2 job function"],
    "job_level": ["position", "job level", "job_level", "seniority", "level", "ocpcode3", "ocpcode3 job level"],
    "job_role": ["job role", "role", "job_role", "job_role__c"],
    "company_size": ["companysize", "company size", "number of employees", "number_of_employees", "employees", "orgemp", "orgemp number of employees"],
    "phone": ["phone", "telephone", "mobile", "cell", "phonemain", "phone main", "phone_1"],
    "website": ["website", "domain", "url", "companyurl", "company url", "web", "company website"],
    "job_title": ["job title", "job_title", "jobtitle", "title", "jobtitletext", "job title text"],
}


def score_header(header: str, alias: str) -> int:
    header_clean = clean_header(header)
    alias_clean = clean_header(alias)
    if not header_clean or not alias_clean:
        return 0
    if header_clean == alias_clean:
        return 100
    if alias_clean in header_clean:
        return 85
    header_tokens = set(header_clean.split())
    alias_tokens = set(alias_clean.split())
    if alias_tokens and alias_tokens.issubset(header_tokens):
        return 80
    return int(fuzz.token_sort_ratio(header_clean, alias_clean))


def detect_columns(df: pd.DataFrame) -> Dict[str, str]:
    result: Dict[str, str] = {}
    used_columns = set()
    for field, aliases in FIELD_ALIASES.items():
        best_col = ""
        best_score = 0
        for col in df.columns:
