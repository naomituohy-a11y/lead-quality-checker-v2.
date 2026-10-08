import io
import unittest
from datetime import date
import pandas as pd
from openpyxl import load_workbook
from checker import detect, domain, company_check, process, phone_check

def excel(df):
    out=io.BytesIO()
    df.to_excel(out,index=False,sheet_name='Leads')
    return out.getvalue()

class Checks(unittest.TestCase):
    def test_headers(self):
        m=detect(pd.DataFrame(columns=['email','companyname','Time in Current Role']))
        self.assertIsNone(m['job_role'])
        self.assertIsNone(m['phone'])
    def test_domain(self):
        self.assertEqual(domain('https://mail.example.co.uk/contact'),'example.co.uk')
        self.assertEqual(domain('a@sub.example.co.uk',True),'example.co.uk')
        self.assertEqual(domain('not an email',True),'')
    def test_relationships(self):
        r={('acme','other.com'):('parent','https://example.com/proof',date.today())}
        self.assertEqual(company_check('Acme','other.com',r)[1],'amber')
        self.assertEqual(company_check('Acme','acme.com',{})[1],'yellow')
        self.assertEqual(company_check('Acme','other.com',{})[1],'yellow')
    def test_email_separate_and_rerun(self):
        df=pd.DataFrame([{'company':'Acme','website':'acme.com','email':'person@old.com','country':'US'}])
        out,results=process(excel(df),'Leads',detect(df))
        self.assertEqual(results.iloc[0]['QA_Email_Website_Status'],'Different domains — review')
        self.assertEqual(results.iloc[0]['QA_Overall_Status'],'REVIEW')
        wb=load_workbook(io.BytesIO(out)); wb['Leads']['C2']='person@acme.com'
        buf=io.BytesIO(); wb.save(buf)
        updated,_=process(buf.getvalue(),'Leads',detect(df))
        ws=load_workbook(io.BytesIO(updated))['Leads']
        headers=[c.value for c in ws[1]]
        self.assertEqual(ws.cell(2,headers.index('QA_Email_Domain')+1).value,'acme.com')
        self.assertEqual(len(headers),len(set(headers)))
    def test_country(self):
        self.assertEqual(phone_check('+442083661177','US'),'Review: phone country differs')

if __name__=='__main__': unittest.main()
